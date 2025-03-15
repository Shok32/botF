import os
import telebot
from telebot.async_telebot import AsyncTeleBot
from telebot.types import InlineKeyboardMarkup, InlineKeyboardButton
from fuzzywuzzy import fuzz
import PyPDF2
import docx
import mammoth
import openpyxl
from odf.opendocument import load as load_odf
from odf.text import P
import requests
import io
import asyncio
import hashlib

# Токен вашего бота от BotFather
TOKEN = '///'

# Путь к локальной папке
LOCAL_PATH = "documents/"

# Ссылка на публичную папку Яндекс.Диска
YANDEX_DISK_PUBLIC_URL = '///'

# Список разрешённых Telegram ID (добавь сюда ID сотрудников)
ALLOWED_IDS = [13131313]  # Пример ID, замени на реальные

# Инициализация асинхронного бота
bot = AsyncTeleBot(TOKEN)

# Кэш документов: {короткий_id: (имя, путь/URL, содержимое)}
documents_cache = {}

# Типы файлов для категорий
DOC_TYPES = ('.doc', '.docx', '.odt', '.txt')
TABLE_TYPES = ('.xls', '.xlsx')
PDF_TYPES = ('.pdf')

# Функция для создания короткого ID
def get_short_id(filename):
    return hashlib.md5(filename.encode('utf-8')).hexdigest()[:8]

# Функция для создания меню
def create_menu():
    markup = InlineKeyboardMarkup(row_width=2)
    buttons = [
        InlineKeyboardButton("📝 Документы", callback_data="category:documents"),
        InlineKeyboardButton("📊 Таблицы", callback_data="category:tables"),
        InlineKeyboardButton("📜 PDF", callback_data="category:pdf"),
        InlineKeyboardButton("🔍 Поиск", callback_data="search"),
        InlineKeyboardButton("📤 Загрузить файл", callback_data="upload")
    ]
    markup.add(*buttons)
    return markup

# Проверка доступа
def check_access(user_id):
    return user_id in ALLOWED_IDS

# Команда /start
@bot.message_handler(commands=['start'])
async def send_welcome(message):
    if not check_access(message.from_user.id):
        await bot.reply_to(message, "❌ Доступ запрещён! Вы не в списке разрешённых пользователей.")
        return
    await bot.reply_to(message, "👋 Привет! Я помогу найти документы по названию или содержимому.\n"
                               "Выбери категорию, начни поиск или загрузи файл:", reply_markup=create_menu())

# Обработка нажатий на кнопки
@bot.callback_query_handler(func=lambda call: True)
async def callback_query(call):
    if not check_access(call.from_user.id):
        await bot.send_message(call.message.chat.id, "❌ Доступ запрещён!")
        return
    
    try:
        await bot.answer_callback_query(call.id)
    except Exception as e:
        print(f"Ошибка ответа на callback: {e}")

    if call.data.startswith("category:"):
        category = call.data.split(":")[1]
        results = search_by_category(category)
        await send_results(call.message, results)
    elif call.data == "search":
        await bot.send_message(call.message.chat.id, "🔎 Введи запрос (название или часть содержимого):")
    elif call.data == "upload":
        await bot.send_message(call.message.chat.id, "📤 Отправь мне файл, и я сохраню его!")
    elif call.data.startswith("file:"):
        short_id = call.data.split(":")[1]
        await send_file(call.message, short_id)

# Обработка текстовых сообщений (поиск)
@bot.message_handler(func=lambda message: True)
async def handle_message(message):
    if not check_access(message.from_user.id):
        await bot.reply_to(message, "❌ Доступ запрещён!")
        return
    
    query = message.text
    results = search_documents(query)
    await send_results(message, results)

# Обработка загруженных файлов
@bot.message_handler(content_types=['document'])
async def handle_docs(message):
    if not check_access(message.from_user.id):
        await bot.reply_to(message, "❌ Доступ запрещён!")
        return
    
    file_info = await bot.get_file(message.document.file_id)
    downloaded_file = await bot.download_file(file_info.file_path)
    file_path = os.path.join(LOCAL_PATH, message.document.file_name)
    
    if not os.path.exists(LOCAL_PATH):
        os.makedirs(LOCAL_PATH)
    
    with open(file_path, 'wb') as f:
        f.write(downloaded_file)
    
    content = read_file(file_path, message.document.file_name)
    short_id = get_short_id(message.document.file_name)
    documents_cache[short_id] = (message.document.file_name, file_path, content)
    print(f"Загружен файл: {message.document.file_name} (ID: {short_id})")
    await bot.reply_to(message, f"✅ Файл '{message.document.file_name}' сохранён!", reply_markup=create_menu())

# Поиск по названию и содержимому
def search_documents(query):
    if not documents_cache:
        load_documents()
    
    results = []
    query = query.lower().strip()
    query_words = query.split()
    print(f"Поиск по запросу: {query}")
    print(f"Текущий кэш: {list(documents_cache.keys())}")
    
    for short_id, (name, source, content) in documents_cache.items():
        name_lower = name.lower()
        if fuzz.partial_ratio(query, name_lower) > 70 or query in name_lower:
            results.append((name, short_id))
        elif content:
            content_lower = content.lower()
            if any(word in content_lower for word in query_words):
                results.append((name, short_id))
    return results[:]

# Поиск по категории
def search_by_category(category):
    if not documents_cache:
        load_documents()
    
    results = []
    if category == "documents":
        file_types = DOC_TYPES
    elif category == "tables":
        file_types = TABLE_TYPES
    elif category == "pdf":
        file_types = PDF_TYPES
    else:
        return results
    
    for short_id, (name, source, content) in documents_cache.items():
        if name.lower().endswith(file_types):
            results.append((name, short_id))
    return results[:]

# Загрузка документов
def load_documents():
    load_local_documents()
    if YANDEX_DISK_PUBLIC_URL:
        load_yandex_disk_documents()

# Локальная загрузка
def load_local_documents():
    if not os.path.exists(LOCAL_PATH):
        os.makedirs(LOCAL_PATH)
    for filename in os.listdir(LOCAL_PATH):
        filepath = os.path.join(LOCAL_PATH, filename)
        if os.path.isfile(filepath):
            content = read_file(filepath, filename)
            short_id = get_short_id(filename)
            documents_cache[short_id] = (filename, filepath, content)
            print(f"Загружен локальный файл: {filename} (ID: {short_id})")

# Загрузка с Яндекс.Диска
def load_yandex_disk_documents():
    api_url = f"https://cloud-api.yandex.net/v1/disk/public/resources?public_key={YANDEX_DISK_PUBLIC_URL}"
    response = requests.get(api_url)
    if response.status_code == 200:
        data = response.json()
        items = data.get('_embedded', {}).get('items', [])
        print(f"Файлы с Яндекс.Диска: {[item['name'] for item in items]}")
        for item in items:
            if item['type'] == 'file':
                name = item['name']
                download_url = item['file']
                content = read_file(download_url, name)
                short_id = get_short_id(name)
                documents_cache[short_id] = (name, download_url, content)
                print(f"Загружен файл с Яндекс.Диска: {name} (ID: {short_id})")
    else:
        print(f"Ошибка загрузки с Яндекс.Диска: {response.status_code} - {response.text}")

# Чтение файла
def read_file(filepath_or_url, filename):
    try:
        if isinstance(filepath_or_url, str) and filepath_or_url.startswith('http'):
            response = requests.get(filepath_or_url)
            file_content = io.BytesIO(response.content)
        else:
            file_content = open(filepath_or_url, 'rb')
        
        if filename.endswith('.txt'):
            return file_content.read().decode('utf-8', errors='ignore')
        elif filename.endswith('.pdf'):
            reader = PyPDF2.PdfReader(file_content)
            text = "".join(page.extract_text() or "" for page in reader.pages)
            return text if text else ""
        elif filename.endswith('.docx'):
            doc = docx.Document(file_content)
            text = " ".join([para.text for para in doc.paragraphs if para.text])
            return text if text else ""
        elif filename.endswith('.doc'):
            result = mammoth.extract_raw_text(file_content)
            text = result.value if result.value else ""
            return text
        elif filename.endswith(('.xls', '.xlsx')):
            wb = openpyxl.load_workbook(file_content)
            text = ""
            for sheet in wb:
                for row in sheet.rows:
                    text += " ".join(str(cell.value or "") for cell in row) + " "
            return text
        elif filename.endswith('.odt'):
            doc = load_odf(file_content)
            text = " ".join(p.text or "" for p in doc.getElementsByType(P))
            return text if text else ""
        return ""
    except Exception as e:
        print(f"Ошибка чтения {filename}: {e}")
        return ""
    finally:
        if not isinstance(file_content, io.BytesIO):
            file_content.close()

# Отправка результатов
async def send_results(message, results):
    if not results:
        await bot.reply_to(message, "❌ Ничего не найдено!", reply_markup=create_menu())
        return
    
    markup = InlineKeyboardMarkup(row_width=1)
    for name, short_id in results:
        markup.add(InlineKeyboardButton(f"📄 {name}", callback_data=f"file:{short_id}"))
    await bot.reply_to(message, "✅ Найденные документы:", reply_markup=markup)
    await bot.send_message(message.chat.id, "Выбери действие:", reply_markup=create_menu())

# Отправка файла
async def send_file(message, short_id):
    if short_id not in documents_cache:
        await bot.reply_to(message, "❌ Файл не найден!", reply_markup=create_menu())
        return
    
    name, source, _ = documents_cache[short_id]
    if source.startswith('http'):  # Яндекс.Диск
        response = requests.get(source)
        file_content = io.BytesIO(response.content)
        await bot.send_document(message.chat.id, file_content, caption=name, reply_markup=create_menu())
    else:  # Локально
        with open(source, 'rb') as f:
            await bot.send_document(message.chat.id, f, caption=name, reply_markup=create_menu())

# Запуск бота
async def main():
    print("Бот запущен...")
    load_documents()
    await bot.polling()

if __name__ == "__main__":
    asyncio.run(main())
